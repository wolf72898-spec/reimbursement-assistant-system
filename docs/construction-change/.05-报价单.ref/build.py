# -*- coding: utf-8 -*-
"""远筑工程 · 工程变更与签证协同 AI 助手项目报价单"""
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value

XL_PRIMARY = xl_color("#4472C4")
XL_LIGHT = xl_color("#D9E2F3")
XL_BORDER = xl_color("#BFBFBF")
XL_SUM = xl_color("#2F5597")
thin = Side(style="thin", color=XL_BORDER)
BORDER_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "报价单"

# ---------- 列宽 ----------
widths = {"A": 6, "B": 12, "C": 26, "D": 40, "E": 8, "F": 8, "G": 12, "H": 14, "I": 18}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# ---------- 标题区（第1行） ----------
ws.merge_cells("A1:I1")
ws["A1"] = "远筑工程 · 工程变更与签证协同 AI 助手项目报价单"
ws["A1"].font = Font(size=16, bold=True, color=XL_PRIMARY)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# ---------- 报价信息区（第2~3行，表单抬头） ----------
ws.merge_cells("A2:E2")
ws["A2"] = "客户（甲方）：远筑工程"
ws.merge_cells("F2:I2")
ws["F2"] = "报价方（乙方）：AI 实施服务商"
ws.merge_cells("A3:E3")
ws["A3"] = "项目名称：工程变更与签证协同数智化试点（首期 3 个月）"
ws.merge_cells("F3:I3")
ws["F3"] = "报价日期：2026-10-30    报价有效期：30 天"
for r in (2, 3):
    for c in range(1, 10):
        ws.cell(row=r, column=c).font = Font(size=10.5)
        ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=XL_LIGHT)

# ---------- 表头（第4行） ----------
headers = ["序号", "类别", "费用项", "内容说明", "单位", "数量", "单价（元）", "小计（元）", "备注"]
for i, h in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=i, value=h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=XL_PRIMARY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrapText=True)
    cell.border = BORDER_ALL
ws.row_dimensions[4].height = 22

# ---------- 明细（第5~13行，锚点：数据起 5、止 13） ----------
items = [
    ("实施服务", "业务挖掘与现场调研", "六角色访谈、机会清单、指标基线与现状流程图", "人天", 10, 4800, "含访谈纪要与收敛表"),
    ("实施服务", "需求定义与方案设计", "需求规格、验收清单、设计方案与实施计划", "人天", 8, 4800, ""),
    ("实施服务", "平台配置与构建集成", "统一变更入口、版本管理、权限中心、消息推送", "人天", 15, 4800, "不含业主侧系统改造"),
    ("实施服务", "AI 能力接入与调优", "图纸比对、变更影响识别、签证单起草提示词调优", "人天", 10, 5600, "含样本验证"),
    ("实施服务", "测试评估与试点支持", "试点上线陪跑、每周复盘、指标对照评估", "人天", 6, 4800, ""),
    ("平台订阅", "协同平台订阅费", "变更/签证协同平台，含 50 用户账号", "月", 3, 8000, "试点期 3 个月"),
    ("AI 用量", "AI 文档处理用量包", "图纸比对、要素提取、起草生成等用量", "万次", 3, 3000, "超出部分按量另计"),
    ("培训交付", "现场培训与操作手册", "管理端与现场端培训各 1 场，含操作手册", "场", 2, 3600, ""),
    ("运维支持", "试点期远程运维支持", "工单响应、故障处理、月度运营报告", "项", 1, 2600, "工作日 5×8 小时"),
]
DATA_START = 5
for idx, (cat, item, desc, unit, qty, price, note) in enumerate(items):
    r = DATA_START + idx
    ws.cell(row=r, column=1, value=idx + 1)
    ws.cell(row=r, column=2, value=cat)
    ws.cell(row=r, column=3, value=item)
    ws.cell(row=r, column=4, value=desc)
    ws.cell(row=r, column=5, value=unit)
    ws.cell(row=r, column=6, value=qty)
    ws.cell(row=r, column=7, value=price)
    ws.cell(row=r, column=8, value=f'=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r})')
    ws.cell(row=r, column=9, value=note)
DATA_END = DATA_START + len(items) - 1  # 13

# 数据区样式
for r in range(DATA_START, DATA_END + 1):
    for c in range(1, 10):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER_ALL
        if c in (1, 5, 6):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c in (7, 8):
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrapText=True)
        if c in (7, 8):
            cell.number_format = "¥#,##0.00"
        if c == 4:
            cell.font = Font(size=10)

# ---------- 合计区（第14~15行） ----------
TOTAL_ROW = DATA_END + 1  # 14
ws.merge_cells(f"A{TOTAL_ROW}:G{TOTAL_ROW}")
ws.cell(row=TOTAL_ROW, column=1, value="合计（含税）")
ws.cell(row=TOTAL_ROW, column=8, value=f"=SUM(H{DATA_START}:H{DATA_END})")
ws.cell(row=TOTAL_ROW, column=9, value="首期试点打包价")
CAP_ROW = TOTAL_ROW + 1  # 15
ws.merge_cells(f"A{CAP_ROW}:I{CAP_ROW}")
ws.cell(row=CAP_ROW, column=1, value="报价总额大写：人民币贰拾捌万陆仟元整（¥286,000.00）")
for r in (TOTAL_ROW, CAP_ROW):
    for c in range(1, 10):
        cell = ws.cell(row=r, column=c)
        cell.fill = PatternFill("solid", fgColor=XL_SUM)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = BORDER_ALL
        cell.alignment = Alignment(horizontal="center" if r == CAP_ROW else "right",
                                   vertical="center")
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
ws.cell(row=TOTAL_ROW, column=8).number_format = "¥#,##0.00"
ws.row_dimensions[TOTAL_ROW].height = 22

# ---------- 备注区（第16~18行） ----------
notes = [
    "备注：1. 以上为首期试点（3 个月）打包价，含税；正式推广阶段按项目数量与用户规模另行报价。",
    "2. 业主侧系统对接改造、电子签章年费、私有化部署等需求不在本报价范围内，另行核算。",
    "3. 付款方式建议：合同签订 50%，试点验收通过 50%。",
]
for i, t in enumerate(notes):
    r = CAP_ROW + 1 + i
    ws.merge_cells(f"A{r}:I{r}")
    ws.cell(row=r, column=1, value=t)
    ws.cell(row=r, column=1).font = Font(size=9.5, color="555555")
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")

# ---------- 签字栏（表单型） ----------
SIGN_ROW = CAP_ROW + len(notes) + 2
ws.merge_cells(f"A{SIGN_ROW}:C{SIGN_ROW}")
ws.merge_cells(f"D{SIGN_ROW}:F{SIGN_ROW}")
ws.merge_cells(f"G{SIGN_ROW}:I{SIGN_ROW}")
ws.cell(row=SIGN_ROW, column=1, value="报价人（乙方）：____________")
ws.cell(row=SIGN_ROW, column=4, value="审核（乙方）：____________")
ws.cell(row=SIGN_ROW, column=7, value="客户确认（甲方）：____________")
ws.cell(row=SIGN_ROW + 1, column=1, value="日期：____________")
ws.cell(row=SIGN_ROW + 1, column=7, value="日期：____________")
for r in (SIGN_ROW, SIGN_ROW + 1):
    for c in range(1, 10):
        ws.cell(row=r, column=c).font = Font(size=10.5)

# ---------- 打印设置 ----------
ws.print_area = f"A1:I{SIGN_ROW + 1}"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True

wb.properties.title = "远筑工程 · 工程变更与签证协同 AI 助手项目报价单"
wb.save(r"D:\aworkbuddy\报销助手智能体系统\docs\construction-change\05-报价单.xlsx")
print("saved ok")
